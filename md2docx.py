# -*- coding: utf-8 -*-
"""Markdown → docx 简易转换（适配本论文稿结构：标题/加粗/表格/列表/引用块）"""
import re
import sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


def set_cn_font(style, size=12, bold=None):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")


def add_runs(p, text):
    """处理 **加粗** 行内标记"""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        else:
            p.add_run(part)


def convert(md_path, docx_path):
    doc = Document()
    set_cn_font(doc.styles["Normal"], 12)
    for i in range(1, 4):
        set_cn_font(doc.styles[f"Heading {i}"], 16 - i, True)

    lines = open(md_path, encoding="utf-8").read().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        # 表格块
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows) + 1, cols=len(header))
            t.style = "Table Grid"
            for j, h in enumerate(header):
                t.rows[0].cells[j].text = h
                for r in t.rows[0].cells[j].paragraphs[0].runs:
                    r.bold = True
            for ri, row in enumerate(rows, 1):
                for j, c in enumerate(row[:len(header)]):
                    t.rows[ri].cells[j].text = c
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            add_runs(p, line[2:].strip())
            for r in p.runs:
                r.italic = True
        elif re.match(r"^\d+\. ", line.strip()):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\. ", "", line.strip()))
        elif line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line.strip()[2:])
        elif line.strip() == "---":
            pass
        elif line.strip():
            p = doc.add_paragraph()
            add_runs(p, line.strip())
        i += 1
    doc.save(docx_path)
    print("saved:", docx_path)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
